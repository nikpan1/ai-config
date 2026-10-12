import csv
import io
import zipfile
from pathlib import PurePosixPath

from defusedxml import ElementTree

PARSER_VERSION = "attachments-v1"
SUPPORTED_ATTACHMENTS = {".txt", ".md", ".pdf", ".docx", ".csv", ".xlsx"}
MAX_EXPANDED_BYTES = 50_000_000


def _text(content: bytes) -> str:
    return content.decode("utf-8-sig")


def _archive(content: bytes):
    archive = zipfile.ZipFile(io.BytesIO(content))
    if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
        archive.close()
        raise ValueError("Expanded attachment exceeds the parser limit")
    return archive


def _docx(content: bytes) -> tuple[list[dict], list[str]]:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    spans = []
    with _archive(content) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        body = root.find(f"{namespace}body")
        for index, element in enumerate(body if body is not None else [], 1):
            if element.tag == f"{namespace}tbl":
                for row_index, row in enumerate(element.findall(f"{namespace}tr"), 1):
                    for column, cell in enumerate(row.findall(f"{namespace}tc"), 1):
                        value = " ".join(node.text or "" for node in cell.iter(f"{namespace}t"))
                        spans.append(
                            {
                                "text": value,
                                "locator": {
                                    "kind": "table_cell",
                                    "table": index,
                                    "row": row_index,
                                    "column": column,
                                },
                            }
                        )
            elif element.tag == f"{namespace}p":
                value = " ".join(node.text or "" for node in element.iter(f"{namespace}t"))
                if value.strip():
                    spans.append(
                        {"text": value, "locator": {"kind": "paragraph", "paragraph": index}}
                    )
        limitations = []
        unsupported = [
            name
            for name in archive.namelist()
            if name.startswith(
                (
                    "word/comments",
                    "word/footnotes",
                    "word/endnotes",
                    "word/header",
                    "word/footer",
                    "word/media/",
                )
            )
        ]
        if unsupported or any(
            node.tag
            in {
                f"{namespace}drawing",
                f"{namespace}altChunk",
                f"{namespace}vMerge",
                f"{namespace}gridSpan",
                f"{namespace}ins",
                f"{namespace}del",
            }
            for node in root.iter()
        ):
            limitations.append(
                "Annotations, drawings, merged cells, headers, or embedded content "
                "require source-attributed interpretation"
            )
    return spans, limitations


def _xlsx(content: bytes) -> tuple[list[dict], list[str]]:
    from openpyxl import load_workbook

    with _archive(content):
        pass
    formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    spans, limitations = [], []
    try:
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cached[sheet.title][cell.coordinate].value
                    text = str(cell.value)
                    if cell.data_type == "f":
                        text += f"; cached value: {value if value is not None else 'unavailable'}"
                        if value is None:
                            limitations.append(
                                f"Missing cached formula value at {sheet.title}!{cell.coordinate}"
                            )
                    spans.append(
                        {
                            "text": text,
                            "locator": {
                                "kind": "sheet_range",
                                "sheet": sheet.title,
                                "range": cell.coordinate,
                            },
                            "formula": str(cell.value) if cell.data_type == "f" else None,
                            "cached_value": str(value) if value is not None else None,
                        }
                    )
        with _archive(content) as archive:
            if any(
                name.startswith(("xl/drawings/", "xl/comments", "xl/charts/"))
                for name in archive.namelist()
            ):
                limitations.append("Workbook drawings, comments, or charts require interpretation")
    finally:
        formulas.close()
        cached.close()
    return spans, limitations


def extract_attachment(path: str, content: bytes) -> dict:
    suffix = PurePosixPath(path.replace("\\", "/")).suffix.casefold()
    result = {
        "path": path,
        "parser_version": PARSER_VERSION,
        "status": "ok",
        "spans": [],
        "limitations": [],
    }
    try:
        if len(content) > MAX_EXPANDED_BYTES:
            raise ValueError("Attachment exceeds the parser limit")
        if suffix in {".txt", ".md"}:
            result["spans"] = [
                {"text": line, "locator": {"kind": "lines", "start": index, "end": index}}
                for index, line in enumerate(_text(content).splitlines(), 1)
                if line.strip()
            ]
        elif suffix == ".csv":
            rows = list(csv.reader(io.StringIO(_text(content))))
            headers = rows[0] if rows else []
            result["spans"] = [
                {
                    "text": (headers[column - 1] if column <= len(headers) else str(column))
                    + f": {value}",
                    "locator": {
                        "kind": "table_cell",
                        "table": 1,
                        "row": row_index,
                        "column": column,
                    },
                }
                for row_index, row in enumerate(rows, 1)
                for column, value in enumerate(row, 1)
                if value.strip()
            ]
        elif suffix == ".docx":
            result["spans"], result["limitations"] = _docx(content)
        elif suffix == ".xlsx":
            result["spans"], result["limitations"] = _xlsx(content)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF requires an accessible source")
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    result["spans"].append(
                        {
                            "text": text,
                            "locator": {
                                "kind": "page",
                                "page": index,
                                "start": 0,
                                "end": len(text),
                            },
                        }
                    )
                else:
                    result["limitations"].append(
                        f"Page {index} has no extractable text; transcription required"
                    )
                if page.get("/Annots") or page.images:
                    result["limitations"].append(
                        f"Page {index} contains annotations or images requiring interpretation"
                    )
        else:
            result.update(
                status="unsupported",
                limitations=[
                    "This format requires source-attributed transcription; no OCR or vision"
                ],
            )
    except Exception as error:
        result.update(
            status="unresolved",
            limitations=[f"Attachment extraction failed: {type(error).__name__}"],
        )
    if result["status"] == "ok" and (result["limitations"] or not result["spans"]):
        result["status"] = "unresolved"
        if not result["spans"]:
            result["limitations"].append("No readable content was extracted")
    return result
